import duckdb
import pandas as pd
import pytest
from openpyxl import load_workbook

from src import c6_rate_file as c6


def link_frames():
    links = pd.DataFrame([
        {"tin": "11-1", "npi": 1, "address_id": "a1", "pair_rank": 1, "link_evidence": "verified_address"},
        {"tin": "11-1", "npi": 1, "address_id": "a2", "pair_rank": 2, "link_evidence": "verified_address"},
        {"tin": "11-1", "npi": 1, "address_id": "a3", "pair_rank": 3, "link_evidence": "npi_address_only"},
        {"tin": "22-2", "npi": 2, "address_id": "b1", "pair_rank": 1, "link_evidence": "npi_address_only"},
        {"tin": "22-2", "npi": 2, "address_id": "b2", "pair_rank": 2, "link_evidence": "npi_address_only"},
        {"tin": "33-3", "npi": 3, "address_id": None, "pair_rank": None, "link_evidence": "no_address"},
        {"tin": "44-4", "npi": 4, "address_id": "d1", "pair_rank": 1, "link_evidence": "npi_address_only"},
    ])
    addr = pd.DataFrame([
        {"npi": 1, "address_id": "a1", "address_line1": "1 A St", "address_line2": "Ste 2", "city": "LA", "state": "CA", "zip": "90001", "verified": True, "in_target_state": True},
        {"npi": 1, "address_id": "a2", "address_line1": "2 B St", "address_line2": "", "city": "LA", "state": "CA", "zip": "90002", "verified": True, "in_target_state": True},
        {"npi": 1, "address_id": "a3", "address_line1": "3 C St", "address_line2": "", "city": "LA", "state": "CA", "zip": "90003", "verified": False, "in_target_state": True},
        {"npi": 2, "address_id": "b1", "address_line1": "4 D St", "address_line2": "", "city": "SF", "state": "CA", "zip": "94101", "verified": False, "in_target_state": True},
        {"npi": 2, "address_id": "b2", "address_line1": "5 E St", "address_line2": "", "city": "SF", "state": "CA", "zip": "94102", "verified": False, "in_target_state": True},
        {"npi": 4, "address_id": "d1", "address_line1": "6 F St", "address_line2": "", "city": "Reno", "state": "NV", "zip": "89501", "verified": False, "in_target_state": False},
    ])
    prov = pd.DataFrame({"npi": [1, 2, 3, 4], "location_status": ["verified", "unverified_multiple_candidates", "no_address", "out_of_state_only"]})
    return links, addr, prov


def test_pair_addresses_keeps_all_verified_else_best():
    p = c6.pair_addresses(*link_frames())
    by = p.groupby("npi")["address_id"].apply(lambda s: sorted(s.dropna()))
    assert by[1] == ["a1", "a2"] and by[2] == ["b1"] and by[3] == [] and by[4] == ["d1"]
    st = dict(zip(p["npi"], p["address_status"]))
    assert st == {1: "verified", 2: "unverified", 3: "no_address", 4: "out_of_state"}
    assert p.loc[p["address_id"] == "a1", "address"].iloc[0] == "1 A St, Ste 2"


def test_excel_split_respects_row_limit(tmp_path, monkeypatch):
    rows = pd.DataFrame([{"npi": npi, "tin": "01-2", "address": "1 A St, LA, CA 90001", "billing_code": "0250", "specialty": "X", "rate": 10.5}
                         for npi in [1, 1, 1, 1, 1, 1, 2, 3, 3]])
    pq = tmp_path / "rates.parquet"
    rows.to_parquet(pq)
    monkeypatch.setattr(c6, "EXCEL_MAX_ROWS", 4)
    sheets = c6.write_excel(duckdb.connect(), pq, tmp_path / "out.xlsx", 2)
    assert sheets == [4, 2, 3]
    wb = load_workbook(tmp_path / "out.xlsx", read_only=True)
    assert wb.sheetnames == ["Rates 1", "Rates 2", "Rates 3"]
    head, first = list(wb["Rates 1"].iter_rows(min_row=1, max_row=2, values_only=True))
    assert list(head) == ["NPI", "TIN", "Address", "Billing Code", "Specialty", "Rate"]
    assert first == (1, "01-2", "1 A St, LA, CA 90001", "0250", "X", 10.5)


def test_deliverable_has_one_row_per_npi_tin_address_code(tmp_path):
    base = {"npi": 1, "tin": "9", "full_address": "1 A St, LA, CA 90001", "billing_code": "300", "specialty": "Lab", "npi_valid": True,
            "negotiation_arrangement": "ffs", "associated_code": None, "billing_code_modifier": None}
    detail = pd.DataFrame([
        {**base, "rate": 100.0},                                   # base contract
        {**base, "rate": 300.0},                                   # second base rate (e.g. another place of service)
        {**base, "rate": 50.0, "associated_code": "36415"},        # not base: has an associated code
        {**base, "billing_code": "301", "rate": 70.0, "billing_code_modifier": "26"},
        {**base, "npi": 0, "npi_valid": False, "rate": 1.0},
    ])
    pq = tmp_path / "detail.parquet"
    detail.to_parquet(pq)
    con = duckdb.connect()
    st = c6.build_deliverable(con, f"read_parquet('{pq}')", tmp_path / "rates.parquet", tmp_path / "rates.csv")
    out = pd.read_csv(tmp_path / "rates.csv", dtype={"tin": str, "billing_code": str})
    assert list(out.columns) == ["npi", "tin", "address", "billing_code", "specialty", "rate"]
    assert out.set_index("billing_code")["rate"].to_dict() == {"300": 100.0, "301": 70.0}
    assert st["rows"] == 2 and st["keys_with_several_context_rates"] == 1


def test_grain_check_rejects_duplicate_rows(tmp_path):
    base = {c: "x" for c in c6.GRAIN} | {"npi": 1, "npi_valid": True}
    pq = tmp_path / "g.parquet"
    pd.DataFrame([base, {**base, "billing_class": "y"}]).to_parquet(pq)
    con = duckdb.connect()
    stats = c6.grain_check(con, f"read_parquet('{pq}')")
    assert stats["duplicate_keys_on_full_grain"] == 0 and stats["keys_npi_tin_address_code_with_several_rows"] == 1
    pd.DataFrame([base, base]).to_parquet(pq)
    with pytest.raises(ValueError, match="duplicate keys"):
        c6.grain_check(con, f"read_parquet('{pq}')")


def test_corroborated_address_status():
    links, addr, prov = link_frames()
    addr["verification_strength"] = ["strict", "corroborated", "", "", "", ""]
    p = c6.pair_addresses(links, addr, prov)
    st = dict(zip(p["address_id"], p["address_status"]))
    assert st["a1"] == "verified" and st["a2"] == "corroborated"
